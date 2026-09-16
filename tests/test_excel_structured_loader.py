"""Real XLSX regressions through the canonical versioned loader."""

from __future__ import annotations
import hashlib
import json
import pytest
from openpyxl import Workbook
from backend.indexing.document_loader import DocumentArtifactMetadata, DocumentLoader


def metadata(version="version-a", tenant="tenant-a"):
    return DocumentArtifactMetadata(
        tenant_id=tenant,
        knowledge_base_id="kb-a",
        document_id="doc-a",
        document_version_id=version,
        acl_tags=("reader",),
        index_version="catalog-v1",
    )


def workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "商品表"
    sheet.append(["产品", "价格", "价格", None, "备注"])
    sheet.append(["苹果", 5, 6, "第一行\n第二行", "## markdown-like"])
    sheet.append(["香蕉", None, 7, "", ""])
    sheet.append([None] * 5)
    second = book.create_sheet("季度汇总")
    second.append(["区域", "Q1", "Q2"])
    second.append(["华东", 120, 150])
    second.append(["华南", 98, "=B3+12"])
    book.save(path)
    book.close()


def load(path, **kwargs):
    return DocumentLoader(**kwargs).load_document(
        str(path), path.name, metadata=metadata()
    )


def test_excel_preserves_original_structured_contract_in_persisted_json(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook(path)
    chunks = load(path)
    parents = [chunk for chunk in chunks if chunk["chunk_level"] == 2]
    leaves = [chunk for chunk in chunks if chunk["chunk_level"] == 3]
    assert len(parents) == 2
    assert len(leaves) == 4
    first = json.loads(leaves[0]["text"])
    assert first["sheet"] == "商品表"
    assert first["row"] == 2
    assert first["values"] == {
        "产品": "苹果",
        "价格": "5",
        "价格_2": "6",
        "column_4": "第一行\n第二行",
        "备注": "## markdown-like",
    }
    table = json.loads(parents[0]["text"])
    assert table["headers"] == list(first["values"])
    assert table["rows"][0]["values"] == first["values"]
    assert len(table["rows"]) == 2
    assert json.loads(leaves[-1]["text"])["values"]["Q2"] == "=B3+12"
    for chunk in chunks:
        assert chunk["tenant_id"] == "tenant-a"
        assert chunk["document_version_id"] == "version-a"
        assert chunk["acl_tags"] == ["reader"]
        assert (
            chunk["content_hash"] == hashlib.sha256(chunk["text"].encode()).hexdigest()
        )


def test_collision_headers_never_lose_values_or_identity(tmp_path):
    path = tmp_path / "collision.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["价格", "价格", "价格_2", None, "column_4", "tenant_id", "acl_tags"])
    sheet.append([1, 2, 3, 4, 5, "attacker", "admin"])
    book.save(path)
    chunks = load(path)
    values = json.loads(chunks[-1]["text"])["values"]
    assert len(values) == 7
    assert list(values.values()) == ["1", "2", "3", "4", "5", "attacker", "admin"]
    assert values["价格_2"] == "3"
    assert all(chunk["tenant_id"] == "tenant-a" for chunk in chunks)
    assert all(chunk["acl_tags"] == ["reader"] for chunk in chunks)


def test_pages_are_bounded_complete_and_version_scoped(tmp_path):
    path = tmp_path / "many.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["编号", "内容"])
    for index in range(61):
        sheet.append([index, "数据" * 8])
    book.save(path)
    loader = DocumentLoader(max_pages=100, max_page_characters=1000)
    chunks = loader.load_document(str(path), path.name, metadata=metadata())
    parents = [chunk for chunk in chunks if chunk["chunk_level"] == 2]
    leaves = [chunk for chunk in chunks if chunk["chunk_level"] == 3]
    assert len(parents) > 3
    assert len(leaves) == 61
    assert all(len(chunk["text"]) <= 1000 for chunk in chunks)
    assert [json.loads(chunk["text"])["row"] for chunk in leaves] == list(range(2, 63))
    parent_ids = {chunk["chunk_id"] for chunk in parents}
    assert all(chunk["parent_chunk_id"] in parent_ids for chunk in leaves)
    for parent in parents:
        assert len(json.loads(parent["text"])["rows"]) <= 25
        assert all(
            child["section_id"] == parent["section_id"]
            for child in leaves
            if child["parent_chunk_id"] == parent["chunk_id"]
        )
    assert chunks == loader.load_document(str(path), path.name, metadata=metadata())
    replacement = loader.load_document(
        str(path), path.name, metadata=metadata("version-b")
    )
    assert not {chunk["chunk_id"] for chunk in chunks} & {
        chunk["chunk_id"] for chunk in replacement
    }


def test_empty_sheets_header_only_and_original_row_numbers(tmp_path):
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    book.active.title = "空表"
    sheet = book.create_sheet("仅表头")
    sheet.append(["A", "B"])
    rows = book.create_sheet("前置空行")
    rows.append([None, None])
    rows.append(["名字", "值"])
    rows.append([None, None])
    rows.append(["样本", "x\u200by"])
    book.save(path)
    chunks = load(path)
    leaves = [
        json.loads(chunk["text"]) for chunk in chunks if chunk["chunk_level"] == 3
    ]
    assert leaves[0]["sheet"] == "仅表头"
    assert leaves[0]["rows"] == []
    assert leaves[1]["row"] == 4
    assert leaves[1]["values"]["值"] == "xy"


def test_limit_failures_are_explicit_and_workbook_closed(tmp_path):
    path = tmp_path / "large.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["value"])
    for _ in range(26):
        sheet.append(["长文本" * 40])
    book.save(path)
    with pytest.raises(ValueError, match="row exceeds"):
        load(path, max_page_characters=100)
    with pytest.raises(ValueError, match="page count"):
        load(path, max_pages=1)
    path.unlink()


def test_xls_retains_upstream_format_dispatch(monkeypatch):
    from langchain_core.documents import Document
    import backend.indexing.document_loader as document_loader

    class XlsLoader:
        def __init__(self, path):
            assert path == "old.xls"

        def load(self):
            return [Document(page_content="legacy Excel format")]

    monkeypatch.setattr(document_loader, "UnstructuredExcelLoader", XlsLoader)
    chunks = DocumentLoader().load_document("old.xls", "old.xls", metadata=metadata())
    assert chunks[-1]["text"] == "legacy Excel format"
    assert chunks[-1]["chunk_level"] == 3
